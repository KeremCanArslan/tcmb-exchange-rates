"""Download official daily exchange rates from the Central Bank of the Republic of Türkiye (TCMB)
for a date range and export them to Excel or CSV.

Usage:
    python tcmb_rates.py --start 2026-09-01 --end 2026-09-30 --codes USD EUR GBP --out rates.xlsx
    python tcmb_rates.py --days 30 --out rates.csv
"""
from __future__ import annotations

import argparse
import csv
import sys
import time
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.tcmb.gov.tr/kurlar/{ym}/{dmy}.xml"
FIELDS = ("ForexBuying", "ForexSelling", "BanknoteBuying", "BanknoteSelling")
HEADERS = {"User-Agent": "tcmb-exchange-rates/1.0 (+https://github.com/KeremCanArslan)"}


def url_for(day: date) -> str:
    return BASE_URL.format(ym=day.strftime("%Y%m"), dmy=day.strftime("%d%m%Y"))


def _num(text: str | None) -> float | None:
    text = (text or "").strip()
    return float(text) if text else None


def parse_xml(content: bytes | str) -> list[dict]:
    root = ET.fromstring(content)
    day = datetime.strptime(root.attrib["Tarih"], "%d.%m.%Y").date()
    rows = []
    for cur in root.iter("Currency"):
        unit = int(cur.findtext("Unit") or 1)
        row = {
            "date": day,
            "code": cur.attrib.get("CurrencyCode") or cur.attrib.get("Kod"),
            "unit": unit,
            "name_tr": (cur.findtext("Isim") or "").strip(),
            "name_en": (cur.findtext("CurrencyName") or "").strip(),
        }
        for f in FIELDS:
            row[f] = _num(cur.findtext(f))
        rows.append(row)
    return rows


class TcmbClient:
    """Fetches one XML bulletin per day. Weekends and public holidays return 404 and are skipped.
    Raw responses are cached on disk so repeated runs do not hit the server again."""

    def __init__(self, session=None, cache_dir: Path | None = Path(".cache"), pause: float = 0.3, retries: int = 2):
        self.session = session or requests.Session()
        self.cache_dir = cache_dir
        self.pause = pause
        self.retries = retries

    def fetch_day(self, day: date) -> list[dict] | None:
        cached = self.cache_dir / f"{day.isoformat()}.xml" if self.cache_dir else None
        if cached and cached.exists():
            return parse_xml(cached.read_bytes())
        for attempt in range(self.retries + 1):
            try:
                resp = self.session.get(url_for(day), headers=HEADERS, timeout=15)
            except requests.RequestException:
                resp = None
            if resp is not None and resp.status_code == 404:
                return None
            if resp is not None and resp.status_code == 200:
                if cached:
                    cached.parent.mkdir(parents=True, exist_ok=True)
                    cached.write_bytes(resp.content)
                return parse_xml(resp.content)
            time.sleep(self.pause * (attempt + 2))
        raise RuntimeError(f"TCMB did not respond for {day.isoformat()}")

    def fetch_range(self, start: date, end: date, codes: list[str] | None = None) -> list[dict]:
        wanted = {c.upper() for c in codes} if codes else None
        rows, day = [], start
        while day <= end:
            if day.weekday() < 5:  # no bulletin on weekends
                found = self.fetch_day(day)
                if found:
                    rows += [r for r in found if wanted is None or r["code"] in wanted]
                if self.pause:
                    time.sleep(self.pause)
            day += timedelta(days=1)
        return rows


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=["date", "code", "unit", "name_tr", "name_en", *FIELDS])
        writer.writeheader()
        for r in rows:
            writer.writerow({**r, "date": r["date"].isoformat()})


def write_xlsx(rows: list[dict], path: Path, field: str = "ForexSelling") -> None:
    wb = Workbook()
    header_font, header_fill = Font(bold=True), PatternFill("solid", fgColor="E7E6E6")

    # Sheet 1: one row per day, one column per currency (selling rate)
    ws = wb.active
    ws.title = "Summary"
    codes = sorted({r["code"] for r in rows}, key=lambda c: (c != "USD", c != "EUR", c))
    days = sorted({r["date"] for r in rows})
    lookup = {(r["date"], r["code"]): r[field] for r in rows}
    ws.append(["Date", *codes])
    for d in days:
        ws.append([d, *[lookup.get((d, c)) for c in codes]])
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "dd.mm.yyyy"
        for cell in row[1:]:
            cell.number_format = "0.0000"
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 12
    if days and codes:
        chart = LineChart()
        chart.title = f"{field} (TRY)"
        chart.height, chart.width = 8, 18
        series = min(len(codes), 3)
        chart.add_data(Reference(ws, min_col=2, max_col=1 + series, min_row=1, max_row=len(days) + 1),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(days) + 1))
        ws.add_chart(chart, f"{get_column_letter(len(codes) + 3)}2")

    # Sheet 2: full data in long format, filterable
    data = wb.create_sheet("Data")
    cols = ["date", "code", "unit", "name_tr", "name_en", *FIELDS]
    data.append(cols)
    for r in rows:
        data.append([r[c] for c in cols])
    for cell in data[1]:
        cell.font, cell.fill = header_font, header_fill
    for row in data.iter_rows(min_row=2):
        row[0].number_format = "dd.mm.yyyy"
    data.auto_filter.ref = data.dimensions
    data.freeze_panes = "A2"
    for i, w in enumerate((12, 8, 6, 26, 26, 14, 14, 16, 16), start=1):
        data.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Export TCMB daily exchange rates to Excel or CSV.")
    ap.add_argument("--start", type=date.fromisoformat)
    ap.add_argument("--end", type=date.fromisoformat, default=date.today())
    ap.add_argument("--days", type=int, default=30, help="used when --start is not given")
    ap.add_argument("--codes", nargs="*", default=["USD", "EUR", "GBP"], help="empty list = all currencies")
    ap.add_argument("--out", type=Path, default=Path("rates.xlsx"))
    args = ap.parse_args(argv)
    start = args.start or args.end - timedelta(days=args.days)
    rows = TcmbClient().fetch_range(start, args.end, args.codes or None)
    if not rows:
        print("No bulletin found in the given range.", file=sys.stderr)
        return 1
    (write_csv if args.out.suffix.lower() == ".csv" else write_xlsx)(rows, args.out)
    print(f"{len(rows)} rows, {len({r['date'] for r in rows})} days -> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
