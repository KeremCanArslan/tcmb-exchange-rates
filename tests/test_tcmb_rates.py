import os
import sys
from datetime import date

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from tcmb_rates import TcmbClient, parse_xml, url_for, write_csv, write_xlsx  # noqa: E402

SAMPLE = """<?xml version="1.0" encoding="UTF-8"?>
<Tarih_Date Tarih="{d}" Date="x" Bulten_No="2026/186">
  <Currency CrossOrder="0" Kod="USD" CurrencyCode="USD">
    <Unit>1</Unit><Isim>ABD DOLARI</Isim><CurrencyName>US DOLLAR</CurrencyName>
    <ForexBuying>48.9699</ForexBuying><ForexSelling>49.0582</ForexSelling>
    <BanknoteBuying>48.9357</BanknoteBuying><BanknoteSelling>49.1318</BanknoteSelling>
  </Currency>
  <Currency CrossOrder="9" Kod="JPY" CurrencyCode="JPY">
    <Unit>100</Unit><Isim>JAPON YENİ</Isim><CurrencyName>JAPENESE YEN</CurrencyName>
    <ForexBuying>32.1</ForexBuying><ForexSelling>32.3</ForexSelling>
    <BanknoteBuying/><BanknoteSelling/>
  </Currency>
</Tarih_Date>"""


class Resp:
    def __init__(self, status, content=b""):
        self.status_code, self.content = status, content


class FakeSession:
    """Serves bulletins for the given days and 404 for everything else."""

    def __init__(self, days):
        self.days = {url_for(d): SAMPLE.format(d=d.strftime("%d.%m.%Y")).encode() for d in days}
        self.calls = 0

    def get(self, url, headers=None, timeout=None):
        self.calls += 1
        return Resp(200, self.days[url]) if url in self.days else Resp(404)


def test_url_format():
    assert url_for(date(2026, 10, 2)) == "https://www.tcmb.gov.tr/kurlar/202610/02102026.xml"


def test_parse_handles_units_and_empty_fields():
    rows = parse_xml(SAMPLE.format(d="02.10.2026"))
    usd, jpy = rows
    assert usd["date"] == date(2026, 10, 2) and usd["ForexSelling"] == 49.0582
    assert jpy["unit"] == 100 and jpy["BanknoteBuying"] is None
    assert jpy["name_tr"] == "JAPON YENİ"


def test_range_skips_weekends_and_holidays(tmp_path):
    # 1-2 Oct 2026 are Thu-Fri, 3-4 are the weekend, 5 is Monday; Monday is treated as a holiday here.
    session = FakeSession([date(2026, 10, 1), date(2026, 10, 2)])
    client = TcmbClient(session=session, cache_dir=tmp_path, pause=0)
    rows = client.fetch_range(date(2026, 10, 1), date(2026, 10, 5), ["usd"])
    assert [r["date"].day for r in rows] == [1, 2]
    assert session.calls == 3  # weekend days are never requested


def test_cache_prevents_second_request(tmp_path):
    session = FakeSession([date(2026, 10, 2)])
    client = TcmbClient(session=session, cache_dir=tmp_path, pause=0)
    client.fetch_day(date(2026, 10, 2))
    client.fetch_day(date(2026, 10, 2))
    assert session.calls == 1


def test_exports(tmp_path):
    session = FakeSession([date(2026, 10, 1), date(2026, 10, 2)])
    rows = TcmbClient(session=session, cache_dir=None, pause=0).fetch_range(date(2026, 10, 1), date(2026, 10, 2))
    write_xlsx(rows, tmp_path / "r.xlsx")
    wb = load_workbook(tmp_path / "r.xlsx")
    summary = wb["Summary"]
    assert [c.value for c in summary[1]] == ["Date", "USD", "JPY"]
    assert summary.max_row == 3 and summary["B2"].value == 49.0582
    assert wb["Data"].max_row == 5
    write_csv(rows, tmp_path / "r.csv")
    assert (tmp_path / "r.csv").read_text(encoding="utf-8-sig").splitlines()[1].startswith("2026-10-01,USD")
